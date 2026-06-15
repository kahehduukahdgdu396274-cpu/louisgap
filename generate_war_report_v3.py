#!/usr/bin/env python3
"""generate_war_report_v3.py — 模板填数（字号+居中+本金）
- shutil.copy2(模板 → 输出) 再 load_workbook 改数据
- 只改 cell.value；新行继承模板 font + alignment（居中）
- 本金从 strategy_eq.json 读取，禁止写死 200
"""
import csv
import json
import os
import re
import shutil
from copy import copy
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

CENTER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DETAIL_FONT_SIZE = 20
DETAIL_STYLE_COL = 11  # 模板「开仓张数」列：微软雅黑 20 号
NEW_EQUITY_COLS = (13, 14)  # 开仓权益 / 平仓后权益
# 20 号字下列宽（Excel 字符单位）；模板 N 列默认 13 过窄
DETAIL_COL_WIDTHS = {
    11: 20.0,  # 开仓张数
    12: 16.0,  # 杠杆
    13: 24.0,  # 开仓权益
    14: 28.0,  # 平仓后权益（五字表头）
}

SRC = "/home/admin/okx_bot/trades.csv"
TPLT = "/home/admin/okx_bot/hermes_template.xlsx"
OUT = "/home/admin/okx_bot/hermes_latest_war_report.xlsx"
JSON_OUT = "/home/admin/okx_bot/hermes_war_report.json"
EQ_PATH = "/home/admin/okx_bot/strategy_eq.json"
VALID_SIDS = {"003", "006", "009", "010", "011", "012", "013", "014"}
CT_VAL = 0.01
DETAIL_SHEETS = VALID_SIDS | {"MCP手动"}
SHARED_POOL_SIDS = frozenset({"009", "010", "011", "012"})

# 与 VPS strategies.py check_entry / check_exit 同步（改逻辑须同改此处）
STRATEGY_HEADERS = {
    "003": {
        "open": (
            "开仓条件：15m | EMA9/21分离≥40点 | 多:EMA9>EMA21且收盘/前收>EMA9且>VWAP | "
            "空:EMA9<EMA21且收盘/前收<EMA9且<VWAP | 方向:双向"
        ),
        "close": "平仓条件：止损 -1% | 持仓≥60s | 反向穿越≥15点(BUFFER)",
    },
    "006": {
        "open": (
            "开仓条件：15m | 多:K钩头上穿D(K<20,ADX≥20) | "
            "空:K钩头下穿D(K>80,价<VWAP) | 方向:双向"
        ),
        "close": "平仓条件：止损 -1% / 止盈 +2% | 持仓≥60s | 极值区K/D反向钩头",
    },
    "009": {
        "open": "开仓条件：15m | 收盘≤BB下轨×1.006 | K<40且K上拐 | RSI<50 | 方向:仅多",
        "close": "平仓条件：止损 -2% | 持仓≥60s | 触及BB中轨 | 浮盈≥0.8%保本",
    },
    "010": {
        "open": "开仓条件：15m | 收盘≤BB下轨×1.008 | K<35且K上拐 | VR>0.9 | 方向:仅多",
        "close": "平仓条件：止损 -2% | 持仓≥60s | 触及BB中轨 | 浮盈≥0.8%保本",
    },
    "011": {
        "open": (
            "开仓条件：15m | 收盘≤BB下轨+0.4×ATR | K<45,RSI<55 | "
            "下影>实体×0.5且ADX<40 | 方向:仅多"
        ),
        "close": "平仓条件：止损 -2% / 止盈 +2% | 持仓≥60s | 触及BB上轨",
    },
    "012": {
        "open": (
            "开仓条件：15m | 价距EMA200≤1.8% | K<50且K上拐 | ADX>18 | 方向:仅多"
        ),
        "close": "平仓条件：止损 -2% | 持仓≥60s | 浮盈≥0.8%且K回落 | 或触及BB上轨",
    },
    "013": {
        "open": (
            "开仓条件：15m | 价在EMA20±1%内 | RSI<40 | 下影>实体×0.3 | ADX≥20 | 方向:仅多"
        ),
        "close": "平仓条件：止损 -1.5% / 止盈 +2.5% | 持仓≥180s | K<D且K>70 或 RSI>75",
    },
    "014": {
        "open": (
            "开仓条件：15m(K线游民) | 价相对EMA200定方向 | StochRSI金/死叉(30/70) | "
            "强势阳/阴线突破前K高/低 | ADX≥15 | 方向:双向"
        ),
        "close": (
            "平仓条件：止损 -1.8% / 止盈 +3.5% | 持仓≥120s | 破信号K低/高 | "
            "或价破入场±1×ATR | 极值区K/D反向"
        ),
    },
}


def calc_pnl_from_prices(direction, entry_px, exit_px, sz):
    try:
        ep, xp, s = float(entry_px), float(exit_px), float(sz)
    except (TypeError, ValueError):
        return None
    if ep <= 0 or xp <= 0 or s <= 0:
        return None
    if direction in ("long", "多"):
        return round((xp - ep) * s * CT_VAL, 4)
    return round((ep - xp) * s * CT_VAL, 4)


def net_return_pct(entry_px, sz, net_pnl):
    """净收益率%：与净盈亏同号。"""
    try:
        ep, s, pnl = float(entry_px), float(sz), float(net_pnl)
    except (TypeError, ValueError):
        return ""
    denom = ep * s * CT_VAL
    if denom <= 0:
        return ""
    return round(pnl / denom * 100, 4)


def dedup_trades(trades):
    """战报生成前去重，避免 cron 重复追加同一 ordId。"""
    seen_closed, seen_open = set(), set()
    out = []
    for t in trades:
        oid = (t.get("ordId") or "").strip()
        st = (t.get("状态") or "").strip()
        if st == "已平仓":
            if not oid or oid in seen_closed:
                continue
            seen_closed.add(oid)
        elif st == "持仓中":
            key = (t.get("策略", ""), oid)
            if not oid or key in seen_open:
                continue
            seen_open.add(key)
        out.append(t)
    return out


def repair_trades_pnl(trades):
    """写 Excel 前最后一道：净盈亏为 0 但有真实价差则补算。"""
    for t in trades:
        if t.get("状态") != "已平仓":
            continue
        try:
            if abs(float(t.get("净盈亏") or 0)) >= 1e-4:
                continue
        except ValueError:
            pass
        est = calc_pnl_from_prices(
            t.get("方向", ""), t.get("入场价", ""), t.get("出场价", ""), t.get("开仓张数", "")
        )
        if est is not None and abs(est) >= 1e-4:
            t["净盈亏"] = str(est)


def repair_trades_return(trades):
    """收益率% 必须与净盈亏同号（扣费后净收益率）。"""
    for t in trades:
        if t.get("状态") != "已平仓":
            continue
        ret = net_return_pct(t.get("入场价"), t.get("开仓张数"), t.get("净盈亏"))
        if ret != "":
            t["收益率%"] = ret


DEFAULT_EQ = {
    "003": 397,
    "006": 200,
    "009": 250,
    "010": 250,
    "011": 250,
    "012": 250,
    "013": 200,
    "014": 200,
}


def short_time(val):
    if not val:
        return ""
    s = str(val).strip()
    m = re.match(r"^(\d{4})-(\d{2}-\d{2})\s+(\d{2}:\d{2})", s)
    if m:
        return f"{m.group(2)} {m.group(3)}"
    return s


def _time_ms(val, prefer_close=False):
    if prefer_close:
        ts = (val.get("出场时间") or val.get("开仓时间") or "").strip()
    else:
        ts = (val.get("开仓时间") or val.get("出场时间") or "").strip()
    if not ts:
        return 0
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
        try:
            return int(datetime.strptime(ts[:19], fmt).timestamp() * 1000)
        except ValueError:
            continue
    return 0


def _sid_baseline(sid, closed_rows, current_cap):
    """从当前权益倒推该策略首笔开仓前余额。"""
    eq = float(current_cap)
    for row in reversed(sorted(closed_rows, key=lambda r: _time_ms(r, prefer_close=True))):
        if row.get("策略") != sid:
            continue
        try:
            eq = round(eq - float(row.get("净盈亏") or 0), 2)
        except (TypeError, ValueError):
            pass
    return max(1.0, eq)


def equity_open_map(sid, closed_rows, current_cap):
    """按时间线回放：每笔开仓时的策略权益（重叠持仓共享当时余额）。"""
    sid_rows = [r for r in closed_rows if r.get("策略") == sid and r.get("状态") == "已平仓"]
    if not sid_rows:
        return {}
    start = _sid_baseline(sid, closed_rows, current_cap)
    events = []
    for row in sid_rows:
        oid = (row.get("ordId") or "").strip()
        events.append((_time_ms(row, prefer_close=False), "open", oid, row))
        events.append((_time_ms(row, prefer_close=True), "close", oid, row))
    events.sort(key=lambda x: (x[0], 0 if x[1] == "open" else 1))
    running = start
    open_map = {}
    for _ts, kind, oid, row in events:
        if kind == "open":
            if oid:
                open_map[oid] = round(running, 2)
        else:
            try:
                running = max(1.0, round(running + float(row.get("净盈亏") or 0), 2))
            except (TypeError, ValueError):
                pass
    return open_map


def load_equity():
    if os.path.exists(EQ_PATH):
        with open(EQ_PATH, encoding="utf-8") as f:
            data = json.load(f)
            return {
                **DEFAULT_EQ,
                **{
                    k: float(v)
                    for k, v in data.items()
                    if k in VALID_SIDS or k == "POOL_009_012"
                },
            }
    return DEFAULT_EQ.copy()


def ensure_detail_sheets(wb):
    """模板缺 010/011/012/014 时从 009 复制，避免明细页空白。"""
    ref = None
    for name in ("009", "003", "006", "013"):
        if name in wb.sheetnames:
            ref = name
            break
    if not ref:
        return
    for sid in sorted(VALID_SIDS):
        if sid not in wb.sheetnames:
            ws = wb.copy_worksheet(wb[ref])
            ws.title = sid


def _capital_label(sid: str, equity: dict) -> str:
    cap = round(float(equity.get(sid, DEFAULT_EQ.get(sid, 200))), 0)
    if sid in SHARED_POOL_SIDS:
        pool = equity.get("POOL_009_012")
        pool_s = f" (池余额≈{round(float(pool), 0):.0f})" if pool else ""
        return f"本金:共享池 ${cap:.0f}{pool_s}"
    return f"本金:独立 ${cap:.0f}"


def apply_detail_column_widths(ws):
    """策略明细页：权益相关列加宽，避免 20 号字挤在一起。"""
    for col, width in DETAIL_COL_WIDTHS.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def apply_strategy_headers(wb, equity: dict):
    """各策略页 R1–R3 写入标题与开平仓条件（与 strategies.py 一致）。"""
    for sid in sorted(VALID_SIDS):
        if sid not in wb.sheetnames:
            continue
        meta = STRATEGY_HEADERS.get(sid)
        if not meta:
            continue
        ws = wb[sid]
        cap = _capital_label(sid, equity)
        ws["A1"].value = f"策略 {sid} · 交易明细"
        ws["A2"].value = f"{meta['open']} | {cap}"
        ws["A3"].value = (
            f"{meta['close']} | 排序:开仓时间 | 开仓权益=开仓时余额 | 平仓后权益=该笔平仓后余额"
        )
        apply_detail_column_widths(ws)
        if ws.max_row >= 5:
            ws.cell(row=5, column=13).value = "开仓权益"
            ws.cell(row=5, column=14).value = "平仓后权益"
            apply_new_equity_column_style(ws, header_row=5, data_row=6)

    if "MCP手动" in wb.sheetnames:
        apply_detail_column_widths(wb["MCP手动"])


def _detail_style_ref(ws, ref_row: int, col: int):
    """新增列无模板样式时，回退到开仓张数列（20 号微软雅黑）。"""
    if col in NEW_EQUITY_COLS or col > DETAIL_STYLE_COL:
        return ws.cell(ref_row, DETAIL_STYLE_COL)
    return ws.cell(ref_row, col)


def apply_new_equity_column_style(ws, header_row: int = 5, data_row: int = 6):
    """开仓权益 / 平仓后权益：表头与数据行统一 20 号字体 + 居中。"""
    ref = ws.cell(data_row, DETAIL_STYLE_COL)
    base_font = copy(ref.font) if ref.font else Font(name="微软雅黑", size=DETAIL_FONT_SIZE)
    if base_font.size != DETAIL_FONT_SIZE:
        base_font = Font(
            name=base_font.name or "微软雅黑",
            size=DETAIL_FONT_SIZE,
            bold=base_font.bold,
            italic=base_font.italic,
            color=base_font.color,
        )
    align = (
        copy(ref.alignment)
        if ref.alignment and ref.alignment.horizontal
        else copy(CENTER_ALIGN)
    )
    for col in NEW_EQUITY_COLS:
        for row in (header_row, data_row):
            cell = ws.cell(row, col)
            cell.font = copy(base_font)
            cell.alignment = copy(align)


def apply_data_row_style(ws, dst_row: int, ref_row: int, max_col: int = 14):
    """新写入行继承模板数据行：字号 + 水平/垂直居中"""
    for c in range(1, max_col + 1):
        ref = _detail_style_ref(ws, ref_row, c)
        dst = ws.cell(dst_row, c)
        if ref.font:
            dst.font = copy(ref.font)
            if c in NEW_EQUITY_COLS and (not dst.font.size or dst.font.size < DETAIL_FONT_SIZE):
                dst.font = Font(
                    name=dst.font.name or "微软雅黑",
                    size=DETAIL_FONT_SIZE,
                    bold=dst.font.bold,
                    italic=dst.font.italic,
                    color=dst.font.color,
                )
        if ref.alignment and ref.alignment.horizontal:
            dst.alignment = copy(ref.alignment)
        else:
            dst.alignment = copy(CENTER_ALIGN)


def validate_trades(trades):
    for t in trades:
        if t.get("状态") != "已平仓":
            continue
        sid = str(t.get("策略", "")).strip()
        if sid not in VALID_SIDS:
            continue
        sz = t.get("开仓张数", "")
        ep = t.get("入场价", "")
        xp = t.get("出场价", "")
        oid = t.get("ordId", "")
        if not sz or str(sz) in ("0", "", "0.0"):
            raise ValueError(f"FAIL-FAST: 策略 {sid} 缺少真实开仓张数 (sz={sz})")
        if not ep or float(ep) <= 0:
            raise ValueError(f"FAIL-FAST: 策略 {sid} 缺少真实入场价 (ep={ep})")
        if not xp or float(xp) <= 0:
            raise ValueError(f"FAIL-FAST: 策略 {sid} 缺少真实出场价 (xp={xp})")
        if not oid:
            raise ValueError(f"FAIL-FAST: 策略 {sid} 缺少真实ordId")


def main():
    if not os.path.exists(SRC):
        print(f"[WAR] trades.csv 不存在 ({SRC})，跳过")
        return
    if not os.path.exists(TPLT):
        print(f"[WAR] 模板文件不存在 ({TPLT})，跳过")
        return

    equity = load_equity()

    with open(SRC, "r", encoding="utf-8") as f:
        trades = dedup_trades(list(csv.DictReader(f)))
    if not trades:
        print("[WAR] trades.csv 为空，生成清零战报")
        generate_empty_war_report()
        return
    repair_trades_pnl(trades)
    repair_trades_return(trades)
    validate_trades(trades)

    def _close_ms(t):
        ts = (t.get("出场时间") or t.get("开仓时间") or "").strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M"):
            try:
                return int(datetime.strptime(ts[:19], fmt).timestamp() * 1000)
            except ValueError:
                continue
        return 0

    closed = [t for t in trades if t.get("状态") == "已平仓"]
    active = [t for t in trades if t.get("状态") == "持仓中"]
    strats = sorted(
        s for s in set(t.get("策略") for t in trades if t.get("策略")) if s in VALID_SIDS
    )

    shutil.copy2(TPLT, OUT)
    wb = load_workbook(OUT)
    ensure_detail_sheets(wb)
    apply_strategy_headers(wb, equity)

    ws = wb["汇总"]
    for r in range(5, ws.max_row + 1):
        for c in range(1, 11):
            ws.cell(row=r, column=c).value = None

    r = 5
    ref_summary = 5
    for idx, sid in enumerate(strats, 1):
        sc = [t for t in closed if t["策略"] == sid]
        sa = [t for t in active if t["策略"] == sid]
        total = len(sc)
        wins = sum(1 for t in sc if float(t.get("净盈亏", 0) or 0) > 0)
        wp = f"{round(wins/total*100,1)}%" if total > 0 else "0.0%"
        tpnl = round(sum(float(t.get("净盈亏", 0) or 0) for t in sc), 2)
        avg = round(tpnl / total, 2) if total > 0 else 0.0
        pos_label = (
            "多"
            if any(t.get("方向") in ("long", "多") for t in sa)
            else "空"
            if any(t.get("方向") in ("short", "空") for t in sa)
            else "-"
        )
        cap = round(float(equity.get(sid, DEFAULT_EQ.get(sid, 200))), 2)
        vals = [
            f"{idx:02d}",
            sid,
            cap,
            total,
            wins,
            total - wins,
            wp,
            tpnl,
            avg,
            pos_label,
        ]
        for ci, v in enumerate(vals, 1):
            ws.cell(row=r, column=ci).value = v
        apply_data_row_style(ws, r, ref_summary, 10)
        r += 1

    ws["A2"].value = (
        f"周期:日内滚动 | 更新:{short_time(datetime.now().strftime('%Y-%m-%d %H:%M'))} | 内部投研"
    )

    for sid in DETAIL_SHEETS:
        if sid not in wb.sheetnames:
            continue
        ws = wb[sid]
        for row in ws.iter_rows(min_row=6, max_row=ws.max_row):
            for cell in row:
                cell.value = None

    for sid in strats:
        if sid not in wb.sheetnames:
            continue
        ws = wb[sid]
        sc = [t for t in closed if t["策略"] == sid]
        sa = [t for t in active if t["策略"] == sid]
        sc.sort(key=lambda t: (_time_ms(t, prefer_close=False), _time_ms(t, prefer_close=True)))
        sa.sort(key=lambda t: _time_ms(t, prefer_close=False))
        cap = float(equity.get(sid, DEFAULT_EQ.get(sid, 200)))
        open_eq = equity_open_map(sid, closed, cap)
        data = sc + sa
        ref_detail = 6
        r = 6
        for di, t in enumerate(data, 1):
            d = t.get("方向", "")
            ep_v = float(t["入场价"]) if t.get("入场价") else ""
            xp_v = float(t["出场价"]) if t.get("出场价") else ""
            pp = t.get("收益率%", "")
            npnl = float(t["净盈亏"]) if t.get("净盈亏") else 0.0
            vals = [
                f"{di:03d}",
                short_time(t.get("开仓时间", "")),
                short_time(t.get("出场时间", "")),
                "多" if d in ("long", "多") else "空" if d in ("short", "空") else d,
                ep_v,
                xp_v,
                pp,
                npnl,
                t.get("状态", ""),
                t.get("持仓时间(分)", ""),
            ]
            extra = []
            if trades and "开仓张数" in trades[0]:
                extra.append(t.get("开仓张数", ""))
            if trades and "杠杆" in trades[0]:
                extra.append(t.get("杠杆", ""))
            oid = (t.get("ordId") or "").strip()
            if (t.get("状态") or "").strip() == "已平仓":
                eq_after = t.get("策略权益金", "")
                eq_before = open_eq.get(oid, "")
                extra.extend([eq_before, eq_after])
            else:
                eq_live = t.get("策略权益金", "") or equity.get(sid, "")
                extra.extend([eq_live, ""])
            if trades and "策略权益金" in trades[0]:
                pass
            vals += extra
            for ci, v in enumerate(vals, 1):
                ws.cell(row=r, column=ci).value = v
            apply_data_row_style(ws, r, ref_detail, 14)
            r += 1

    wb.save(OUT)
    shutil.copy2(OUT, os.path.join(os.path.dirname(OUT), "hermes_war_report.xlsx"))
    export_war_json(strats, closed, active, equity, trades)
    print(f"OK: {OUT}")


def export_war_json(strats, closed, active, equity, trades):
    """Mirror war report summary for AI/scripts (Louis keeps xlsx)."""
    summary = []
    for sid in strats:
        sc = [t for t in closed if t["策略"] == sid]
        sa = [t for t in active if t["策略"] == sid]
        total = len(sc)
        wins = sum(1 for t in sc if float(t.get("净盈亏", 0) or 0) > 0)
        tpnl = round(sum(float(t.get("净盈亏", 0) or 0) for t in sc), 2)
        summary.append({
            "strategy_id": sid,
            "capital": round(float(equity.get(sid, DEFAULT_EQ.get(sid, 200))), 2),
            "closed_trades": total,
            "wins": wins,
            "win_rate_pct": round(wins / total * 100, 1) if total > 0 else 0.0,
            "total_pnl": tpnl,
            "open_positions": len(sa),
            "position_direction": (
                "long" if any(t.get("方向") in ("long", "多") for t in sa)
                else "short" if any(t.get("方向") in ("short", "空") for t in sa)
                else None
            ),
        })
    payload = {
        "schema_version": 1,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "summary": summary,
        "total_closed": len(closed),
        "total_open": len(active),
        "total_pnl": round(sum(float(t.get("净盈亏", 0) or 0) for t in closed), 2),
        "trades_count": len(trades),
    }
    tmp = JSON_OUT + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    os.replace(tmp, JSON_OUT)
    print(f"OK: {JSON_OUT}")


def generate_empty_war_report():
    """清零战报：清空模板示例数据，写入 8 策略零成交汇总（本金读 strategy_eq）。"""
    if not os.path.exists(TPLT):
        print(f"[WAR] 模板文件不存在 ({TPLT})，跳过")
        return

    equity = load_equity()
    strats = sorted(VALID_SIDS)

    shutil.copy2(TPLT, OUT)
    wb = load_workbook(OUT)
    ensure_detail_sheets(wb)
    apply_strategy_headers(wb, equity)

    ws = wb["汇总"]
    for r in range(5, ws.max_row + 1):
        for c in range(1, 11):
            ws.cell(row=r, column=c).value = None

    r = 5
    ref_summary = 5
    for idx, sid in enumerate(strats, 1):
        cap = round(float(equity.get(sid, DEFAULT_EQ.get(sid, 200))), 2)
        vals = [f"{idx:02d}", sid, cap, 0, 0, 0, "0.0%", 0, 0, "-"]
        for ci, v in enumerate(vals, 1):
            ws.cell(row=r, column=ci).value = v
        apply_data_row_style(ws, r, ref_summary, 10)
        r += 1

    ws["A2"].value = (
        f"周期:日内滚动 | 更新:{short_time(datetime.now().strftime('%Y-%m-%d %H:%M'))} | 内部投研 | 已清零"
    )

    for sid in DETAIL_SHEETS:
        if sid not in wb.sheetnames:
            continue
        ws = wb[sid]
        for row in ws.iter_rows(min_row=6, max_row=ws.max_row):
            for cell in row:
                cell.value = None

    wb.save(OUT)
    shutil.copy2(OUT, os.path.join(os.path.dirname(OUT), "hermes_war_report.xlsx"))
    export_war_json(strats, [], [], equity, [])
    print(f"OK: empty war report -> {OUT}")


if __name__ == "__main__":
    import sys
    if "--empty" in sys.argv:
        generate_empty_war_report()
    else:
        main()
